"""Test cases for payroll."""

from io import BytesIO

from django.test import SimpleTestCase, TestCase
from openpyxl import Workbook, load_workbook

from employee.models import Employee
from payroll.services.payment_file import (
    fill_payment_file,
    payment_file_has_existing_values,
)
from payroll.views.component_views import _sync_employee_payroll_identity


class EmployeePayrollIdentitySyncTests(TestCase):
    @staticmethod
    def _employee(index, **identity):
        employee = Employee(
            employee_first_name=f"Dipendente {index}",
            email=f"dipendente{index}@example.com",
            phone=f"000{index}",
            **identity,
        )
        Employee.objects.bulk_create([employee])
        return employee

    def test_adds_missing_payroll_code_using_badge(self):
        employee = self._employee(1, badge_id="B001")

        result = _sync_employee_payroll_identity("M001", "B001")

        employee.refresh_from_db()
        self.assertEqual(result, "updated")
        self.assertEqual(employee.codice_paghe, "M001")

    def test_adds_missing_badge_using_payroll_code(self):
        employee = self._employee(2, codice_paghe="M002")

        result = _sync_employee_payroll_identity("M002", "B002")

        employee.refresh_from_db()
        self.assertEqual(result, "updated")
        self.assertEqual(employee.badge_id, "B002")

    def test_replaces_stale_payroll_code_using_badge(self):
        employee = self._employee(5, badge_id="B005", codice_paghe="OLD-M005")

        result = _sync_employee_payroll_identity("M005", "B005")

        employee.refresh_from_db()
        self.assertEqual(result, "updated")
        self.assertEqual(employee.codice_paghe, "M005")

    def test_replaces_stale_badge_using_payroll_code(self):
        employee = self._employee(6, badge_id="OLD-B006", codice_paghe="M006")

        result = _sync_employee_payroll_identity("M006", "B006")

        employee.refresh_from_db()
        self.assertEqual(result, "updated")
        self.assertEqual(employee.badge_id, "B006")

    def test_does_not_overwrite_two_existing_associations(self):
        badge_employee = self._employee(3, badge_id="B003")
        payroll_employee = self._employee(4, codice_paghe="M004")

        result = _sync_employee_payroll_identity("M004", "B003")

        badge_employee.refresh_from_db()
        payroll_employee.refresh_from_db()
        self.assertEqual(result, "conflict")
        self.assertIsNone(badge_employee.codice_paghe)
        self.assertIsNone(payroll_employee.badge_id)

    def test_does_not_create_an_employee_without_a_match(self):
        result = _sync_employee_payroll_identity("M999", "B999")

        self.assertEqual(result, "not_found")
        self.assertFalse(Employee.objects.filter(badge_id="B999").exists())


class PaymentFileServiceTests(SimpleTestCase):
    @staticmethod
    def _workbook_bytes(column_l="acconto esistente", column_m=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Pagamenti"
        sheet["H2"] = "B001"
        sheet["L2"] = column_l
        sheet["M2"] = column_m
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_fill_updates_only_column_m(self):
        content = self._workbook_bytes(column_m="vecchio netto")

        result, filled, missing = fill_payment_file(
            content, "pagamenti.xlsx", "Pagamenti", {"B001": 1234.56}
        )

        workbook = load_workbook(BytesIO(result), data_only=False)
        sheet = workbook["Pagamenti"]
        self.assertEqual(sheet["L2"].value, "acconto esistente")
        self.assertEqual(sheet["M2"].value, 1234.56)
        self.assertEqual(filled, 1)
        self.assertEqual(missing, 0)
        workbook.close()

    def test_existing_value_check_ignores_column_l(self):
        only_l = self._workbook_bytes(column_m=None)
        with_m = self._workbook_bytes(column_m="netto esistente")

        self.assertFalse(
            payment_file_has_existing_values(
                only_l, "pagamenti.xlsx", "Pagamenti", {"B001"}
            )
        )
        self.assertTrue(
            payment_file_has_existing_values(
                with_m, "pagamenti.xlsx", "Pagamenti", {"B001"}
            )
        )
