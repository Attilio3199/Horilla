"""Conversione dei prospetti PDF dei fondi integrativi in file Excel."""

import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


class IntegrativeFundsPdfError(ValueError):
    """Il file caricato non contiene un prospetto PDF leggibile."""


COMPLETE_HEADERS = [
    "matr.", "cognome nome", "data iscr.", "tab.f.", "iscriz.D.L.",
    "iscriz.Lav.", "v.D.L.base", "v.D.L.agg.", "v.D.L.Welf.",
    "v.D.L.Welf.c.", "v.Lav.base", "v.Lav.agg.", "v.TFR parz.",
    "v.TFR tot.", "v.TFR int.", "vers.totale",
]
REDUCED_COLUMNS = (1, 6, 10, 13, 15)
NUMERIC_COLUMNS = tuple(range(4, 16))
DATE_RE = re.compile(r"^\d{1,2}\.\d{2}\.\d{4}$")
ITALIAN_NUMBER_RE = re.compile(r"[-+]?\d{1,3}(?:\.\d{3})*,\d{2}|[-+]?\d+,\d{2}")

# Coordinate x delle 16 colonne nel prospetto DW050612. I valori sono letti
# nelle fasce, non separati da spazi: questo evita lo sfalsamento dei decimali.
COLUMN_BOUNDS = (0, 45, 145, 195, 229, 280, 333, 384, 431, 474, 535, 585, 632, 686, 737, 784)


def _pdf_lines(words):
    """Raggruppa le parole del PDF in righe visive ordinate."""
    lines = defaultdict(list)
    for x0, y0, x1, y1, text, *_ in words:
        if text.strip():
            lines[round(y0 / 3) * 3].append((x0, text.strip()))
    return [sorted(words) for _, words in sorted(lines.items())]


def _is_employee_line(line):
    """Distingue una riga dipendente da titoli, intestazioni e totali."""
    has_matricola = any(x < 45 and text.isdigit() for x, text in line)
    has_date = any(145 <= x < 195 and DATE_RE.match(text) for x, text in line)
    return has_matricola and has_date


def _value_at_column(x, text, row):
    for index, start in reversed(list(enumerate(COLUMN_BOUNDS))):
        if x >= start:
            row[index] = f"{row[index]} {text}".strip()
            return


def _italian_decimal(value):
    """Converte il primo importo italiano trovato in un vero numero Excel."""
    match = ITALIAN_NUMBER_RE.search(str(value or ""))
    if not match:
        return None
    try:
        return Decimal(match.group(0).replace(".", "").replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _normalize_numeric_columns(rows):
    """Rimuove note testuali e converte gli importi in Decimal calcolabili."""
    for row in rows:
        for index in NUMERIC_COLUMNS:
            row[index] = _italian_decimal(row[index])


def _extract_page_rows(page):
    """Estrae dipendenti e totale di una pagina del prospetto."""
    rows = []
    current_row = None

    for line in _pdf_lines(page.get_text("words", sort=True)):
        if _is_employee_line(line):
            current_row = [""] * len(COMPLETE_HEADERS)
            date_found = False
            for x, text in line:
                # La data e' l'unico valore ammesso in ``data iscr.``. I
                # numeri immediatamente dopo appartengono sempre a ``tab.f.``
                # anche quando il PDF li sposta leggermente a sinistra.
                if 145 <= x < 229 and DATE_RE.match(text):
                    current_row[2] = text
                    date_found = True
                    continue
                if date_found and x < 229:
                    current_row[3] = f"{current_row[3]} {text}".strip()
                    continue
                _value_at_column(x, text, current_row)
            rows.append(current_row)
            continue

        if any(text == "TOT.C.FONDO" for _, text in line):
            total_row = [""] * len(COMPLETE_HEADERS)
            for x, text in line:
                _value_at_column(x, text, total_row)
            rows.append(total_row)
            current_row = None
            continue

        # La riga inferiore contiene codice fiscale e descrizione (es.
        # "mensile"). Il codice non fa parte delle 16 colonne; la
        # descrizione viene mantenuta nella cella iscriz.D.L.
        if current_row:
            descriptions = [text for x, text in line if 220 <= x < 280]
            has_fiscal_code = any(45 <= x < 145 for x, _ in line)
            if descriptions and has_fiscal_code:
                current_row[4] = "\n".join(filter(None, [current_row[4], " ".join(descriptions)]))

    return rows


def _fund_code(page):
    match = re.search(r"Cod\.fondo:\s*(.*?)\s+Cod\.Albo:", page.get_text("text", sort=True))
    if not match:
        raise IntegrativeFundsPdfError("Cod.fondo non trovato in una pagina del PDF.")
    return match.group(1).strip()


def _safe_sheet_title(code, existing_titles):
    title = re.sub(r"[\\\\/*?:\[\]]", "-", code)[:31] or "Fondo"
    base_title = title
    suffix = 2
    while title in existing_titles:
        suffix_text = f" ({suffix})"
        title = f"{base_title[:31 - len(suffix_text)]}{suffix_text}"
        suffix += 1
    return title


def _table_name(code, used_names):
    """Crea un nome tabella Excel valido, conservando il codice nel titolo visibile."""
    name = re.sub(r"[^A-Za-z0-9_]", "_", code)
    if not name or name[0].isdigit():
        name = f"Fondo_{name}"
    name = name[:240]
    original_name = name
    suffix = 2
    while name in used_names:
        name = f"{original_name}_{suffix}"
        suffix += 1
    return name


def _write_worksheet(workbook, rows, export_format, fund_code):
    sheet = workbook.create_sheet(_safe_sheet_title(fund_code, workbook.sheetnames))
    if export_format == "reduced":
        headers = [COMPLETE_HEADERS[index] for index in REDUCED_COLUMNS]
        reduced_rows = []
        for row in rows:
            reduced_row = [row[index] for index in REDUCED_COLUMNS]
            # Nel ridotto la prima colonna e' il nome: per il totale diventa
            # l'etichetta TOT.C.FONDO, altrimenti la riga non sarebbe visibile.
            if str(row[0]).startswith("TOT.C.FONDO"):
                reduced_row[0] = row[0]
            reduced_rows.append(reduced_row)
        rows = reduced_rows
    else:
        headers = COMPLETE_HEADERS

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text="\n" in str(value))
            if (
                (export_format == "complete" and column - 1 in NUMERIC_COLUMNS)
                or (export_format == "reduced" and column > 1)
            ):
                cell.number_format = "#,##0.00"
        if str(row[0]).startswith("TOT.C.FONDO"):
            for cell in sheet[row_number]:
                cell.font = Font(bold=True)

    for column_cells in sheet.columns:
        width = max((len(max(str(cell.value or "").split("\n"), key=len)) for cell in column_cells), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 2, 13), 28)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_reduced_worksheet(workbook, fund_groups):
    """Scrive tutti i fondi in un solo foglio, come tabelle Excel distinte."""
    sheet = workbook.create_sheet("Fondi ridotto")
    headers = [COMPLETE_HEADERS[index] for index in REDUCED_COLUMNS]
    header_fill = PatternFill("solid", fgColor="1F4E78")
    current_row = 1
    used_table_names = set()

    for group in fund_groups:
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
        title_cell = sheet.cell(row=current_row, column=1, value=group["code"])
        title_cell.font = Font(bold=True, size=14)
        current_row += 1
        header_row = current_row

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=current_row, column=column, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        current_row += 1

        for row in group["rows"]:
            reduced_row = [row[index] for index in REDUCED_COLUMNS]
            if str(row[0]).startswith("TOT.C.FONDO"):
                reduced_row[0] = row[0]
            for column, value in enumerate(reduced_row, start=1):
                cell = sheet.cell(row=current_row, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text="\n" in str(value))
                if column > 1:
                    cell.number_format = "#,##0.00"
                if str(reduced_row[0]).startswith("TOT.C.FONDO"):
                    cell.font = Font(bold=True)
            current_row += 1

        # Il titolo e' il Cod.fondo leggibile nel foglio; alla tabella Excel
        # viene assegnato un identificatore valido equivalente (Excel vieta spazi).
        table = Table(
            displayName=_table_name(group["code"], used_table_names),
            ref=f"A{header_row}:E{current_row - 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=False, showColumnStripes=False,
        )
        sheet.add_table(table)
        used_table_names.add(table.displayName)
        current_row += 2

    for column_cells in sheet.columns:
        width = max((len(max(str(cell.value or "").split("\n"), key=len)) for cell in column_cells), default=0)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(width + 2, 15), 30)
    sheet.freeze_panes = "A2"


def convert_integrative_funds_pdf(pdf_content, export_format="complete"):
    """Restituisce un XLSX completo o ridotto, elaborando il PDF in memoria."""
    if export_format not in {"complete", "reduced"}:
        raise IntegrativeFundsPdfError("Formato di esportazione non valido.")
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - dipendenza dichiarata
        raise IntegrativeFundsPdfError("Il lettore PDF non e' disponibile.") from exc

    try:
        document = fitz.open(stream=pdf_content, filetype="pdf")
        if document.is_encrypted and not document.authenticate(""):
            raise IntegrativeFundsPdfError("Il PDF e' protetto da password.")
        if document.page_count == 0:
            raise IntegrativeFundsPdfError("Il PDF non contiene pagine.")
        fund_groups = []
        for page in document:
            fund_code = _fund_code(page)
            page_rows = _extract_page_rows(page)
            _normalize_numeric_columns(page_rows)
            if fund_groups and fund_groups[-1]["code"] == fund_code:
                fund_groups[-1]["rows"].extend(page_rows)
            else:
                fund_groups.append({"code": fund_code, "rows": page_rows})
        document.close()
    except IntegrativeFundsPdfError:
        raise
    except Exception as exc:
        raise IntegrativeFundsPdfError("Il file caricato non e' un PDF valido.") from exc

    if not any(group["rows"] for group in fund_groups):
        raise IntegrativeFundsPdfError(
            "Non e' stato possibile estrarre dipendenti dal PDF. Verificare che non sia una scansione."
        )

    workbook = Workbook()
    workbook.remove(workbook.active)
    if export_format == "reduced":
        _write_reduced_worksheet(workbook, fund_groups)
    else:
        for group in fund_groups:
            _write_worksheet(workbook, group["rows"], export_format, group["code"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
