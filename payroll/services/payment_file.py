"""Lettura e compilazione dei file di pagamento dipendenti XLS/XLSX."""

import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

EXCEL_CODE_COLUMN = 8  # H
NET_PAY_COLUMN = 13  # M
EURO_NUMBER_FORMAT = '€ #,##0.00'


class PaymentFileError(ValueError):
    """Il file di pagamento non puo' essere letto o aggiornato."""


def _extension(filename):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return ".xlsx"
    if name.endswith(".xls"):
        return ".xls"
    raise PaymentFileError("Il file deve essere in formato XLS o XLSX.")


def workbook_sheet_names(content, filename):
    """Restituisce i fogli, riparando gli XLSX tolleranti solo di Excel."""
    _, names = prepare_payment_workbook(content, filename)
    return names


def prepare_payment_workbook(content, filename):
    """Valida il file e restituisce contenuto utilizzabile e nomi dei fogli.

    Alcuni file prodotti da gestionali vengono aperti da Excel ma contengono XML
    non strettamente conforme.  openpyxl li rifiuta; LibreOffice li risalva in
    un XLSX valido prima che il file venga memorizzato nella sessione.
    """
    extension = _extension(filename)
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            names = workbook.sheetnames
            workbook.close()
            return content, names
        import xlrd
        return content, xlrd.open_workbook(file_contents=content, on_demand=True).sheet_names()
    except Exception as exc:
        if extension != ".xlsx":
            raise PaymentFileError("Impossibile leggere i fogli del file Excel.") from exc

    try:
        repaired_content = _repair_xlsx(content)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(repaired_content), read_only=True, data_only=False)
        names = workbook.sheetnames
        workbook.close()
        return repaired_content, names
    except Exception as exc:
        raise PaymentFileError("Impossibile leggere i fogli del file Excel.") from exc


def _employee_code(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    code = str(value).strip()
    return code[:-2] if code.endswith(".0") and code[:-2].isdigit() else code


def _has_value(value):
    return value is not None and value != ""


def payment_file_has_existing_values(content, filename, sheet_name, known_codes):
    """Controlla M solo per Badge ID H realmente presenti in Horilla."""
    extension = _extension(filename)
    try:
        if extension == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            if sheet_name not in workbook.sheetnames:
                raise PaymentFileError("Il foglio selezionato non e' presente nel file.")
            sheet = workbook[sheet_name]
            # In read-only mode, Worksheet.cell(row=...) reparses the XML stream
            # on every call.  Calling it three times per row makes large payroll
            # files effectively quadratic.  iter_rows consumes the sheet once.
            found = False
            for cells in sheet.iter_rows(
                min_col=EXCEL_CODE_COLUMN,
                max_col=NET_PAY_COLUMN,
                values_only=True,
            ):
                code = _employee_code(cells[0])
                net_pay = cells[NET_PAY_COLUMN - EXCEL_CODE_COLUMN]
                if code in known_codes and _has_value(net_pay):
                    found = True
                    break
            workbook.close()
            return found
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
        if sheet_name not in workbook.sheet_names():
            raise PaymentFileError("Il foglio selezionato non e' presente nel file.")
        sheet = workbook.sheet_by_name(sheet_name)
        return any(
            _employee_code(sheet.cell_value(row, EXCEL_CODE_COLUMN - 1)) in known_codes
            and _has_value(sheet.cell_value(row, NET_PAY_COLUMN - 1))
            for row in range(sheet.nrows)
        )
    except PaymentFileError:
        raise
    except Exception as exc:
        raise PaymentFileError("Impossibile controllare le celle gia' compilate.") from exc


def _fill_xlsx(content, sheet_name, values, write_mode):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(content), data_only=False)
    except Exception as exc:
        raise PaymentFileError("Impossibile aggiornare il file XLSX.") from exc
    if sheet_name not in workbook.sheetnames:
        raise PaymentFileError("Il foglio selezionato non e' presente nel file.")
    sheet = workbook[sheet_name]
    filled = missing = 0

    # Le righe di riepilogo/calcolo non hanno un codice dipendente in H ma
    # spesso contengono formule in M. Non fanno parte dell'import e devono
    # restare inalterate anche dopo il salvataggio del workbook.
    protected_formulas = {}
    for cells in sheet.iter_rows(
        min_col=EXCEL_CODE_COLUMN,
        max_col=NET_PAY_COLUMN,
    ):
        code = _employee_code(cells[0].value)
        net_pay_cell = cells[NET_PAY_COLUMN - EXCEL_CODE_COLUMN]
        if not code:
            if isinstance(net_pay_cell.value, str) and net_pay_cell.value.startswith("="):
                protected_formulas[net_pay_cell.coordinate] = net_pay_cell.value
            continue
        if code not in values:
            # Nessuna modifica alla riga se il codice in H non ha dati associati.
            missing += 1
            continue
        if write_mode == "blanks_only" and _has_value(net_pay_cell.value):
            continue
        # L'assegnazione mantiene bordi, colori e commenti della cella.
        net_pay_cell.value = values[code]
        net_pay_cell.number_format = EURO_NUMBER_FORMAT
        filled += 1

    # openpyxl normalmente preserva queste formule; il ripristino esplicito
    # rende invarianti le righe che non hanno alcun valore in colonna H.
    for coordinate, formula in protected_formulas.items():
        sheet[coordinate].value = formula

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), filled, missing


def _soffice_convert(source_path, output_format, output_directory, profile_directory):
    """Converte un XLS senza appiattire le formule in valori statici."""
    try:
        result = subprocess.run(
            ["soffice", "--headless", f"-env:UserInstallation={profile_directory.as_uri()}",
             "--convert-to", output_format, "--outdir", str(output_directory), str(source_path)],
            check=False, capture_output=True, text=True, timeout=30,
        )
    except subprocess.TimeoutExpired as exc:
        raise PaymentFileError(
            "La conversione del file Excel ha superato 30 secondi. "
            "Salvare il file come vero XLSX e riprovare."
        ) from exc
    except OSError as exc:
        raise PaymentFileError("Impossibile avviare la conversione del file XLS.") from exc
    if result.returncode != 0:
        raise PaymentFileError("Impossibile convertire il file XLS senza perdere le formule.")


def _repair_xlsx(content):
    """Riscrive un XLSX non conforme senza toccare le celle applicative."""
    with tempfile.TemporaryDirectory(prefix="horilla_xlsx_repair_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        source_path = temp_dir / "source.xlsx"
        output_dir = temp_dir / "output"
        output_dir.mkdir()
        source_path.write_bytes(content)
        _soffice_convert(source_path, "xlsx", output_dir, temp_dir / "profile")
        repaired_path = output_dir / "source.xlsx"
        if not repaired_path.exists():
            raise PaymentFileError("La riparazione del file XLSX non e' riuscita.")
        return repaired_path.read_bytes()


def _fill_xls(content, sheet_name, values, write_mode):
    """Aggiorna un XLS preservando formule, formattazione e tutti i fogli."""
    with tempfile.TemporaryDirectory(prefix="horilla_xls_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        source_path = temp_dir / "source.xls"
        xlsx_dir, xls_dir = temp_dir / "xlsx", temp_dir / "xls"
        xlsx_dir.mkdir()
        xls_dir.mkdir()
        source_path.write_bytes(content)
        _soffice_convert(source_path, "xlsx", xlsx_dir, temp_dir / "profile_xlsx")
        converted_xlsx = xlsx_dir / "source.xlsx"
        if not converted_xlsx.exists():
            raise PaymentFileError("La conversione del file XLS non ha prodotto un XLSX valido.")
        updated_xlsx, filled, missing = _fill_xlsx(
            converted_xlsx.read_bytes(), sheet_name, values, write_mode
        )
        updated_xlsx_path = temp_dir / "compiled.xlsx"
        updated_xlsx_path.write_bytes(updated_xlsx)
        _soffice_convert(updated_xlsx_path, "xls:MS Excel 97", xls_dir, temp_dir / "profile_xls")
        output_xls = xls_dir / "compiled.xls"
        if not output_xls.exists():
            raise PaymentFileError("La conversione finale in XLS non e' riuscita.")
        return output_xls.read_bytes(), filled, missing


def fill_payment_file(content, filename, sheet_name, values, write_mode="overwrite"):
    """Compila M (netto), lasciando L invariata e restituendo il formato di input."""
    extension = _extension(filename)
    if write_mode not in {"overwrite", "blanks_only"}:
        raise PaymentFileError("Modalita' di scrittura non valida.")
    return _fill_xlsx(content, sheet_name, values, write_mode) if extension == ".xlsx" else _fill_xls(content, sheet_name, values, write_mode)
